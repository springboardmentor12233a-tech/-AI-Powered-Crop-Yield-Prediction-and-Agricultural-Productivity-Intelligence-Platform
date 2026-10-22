# -*- coding: utf-8 -*-
"""
AgriYield AI Report Generator - Common Utilities & Styling
"""

from pathlib import Path
from PIL import Image

import docx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docx.shared import Inches, Pt, RGBColor

ROOT_DIR = Path("/home/ramakrishna/Desktop/Projects/AgriYield_AI")
IMAGES_DIR = ROOT_DIR / "Images"
OUTPUT_FILE = ROOT_DIR / "AgriYield_AI_Final_Project_Report_OVS_VARUN.docx"

# Color Palette Constants
COLOR_PRIMARY_HEX = "0B6E4F"       # Deep Forest / Emerald Green
COLOR_SECONDARY_HEX = "1F4E3D"     # Dark Slate Green
COLOR_ACCENT_HEX = "D97706"        # Warm Amber / Gold
COLOR_TEXT_DARK_HEX = "1F2937"     # Dark Charcoal
COLOR_MUTED_HEX = "4B5563"         # Slate Gray
COLOR_LIGHT_BG_HEX = "F0FDF4"      # Soft Mint Green
COLOR_ALT_ROW_HEX = "F8FAFC"       # Very Light Slate
COLOR_BORDER_HEX = "CBD5E1"        # Subtle Border Gray

COLOR_PRIMARY = RGBColor(11, 110, 79)
COLOR_SECONDARY = RGBColor(31, 78, 61)
COLOR_ACCENT = RGBColor(217, 119, 6)
COLOR_TEXT_DARK = RGBColor(31, 41, 55)
COLOR_MUTED = RGBColor(75, 85, 99)


def set_cell_shading(cell, color_hex):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tc_pr.append(tc_mar)


def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders_elm = OxmlElement('w:tcBorders')
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for border_name, border_style in borders.items():
        if border_style:
            val, sz, space, color = border_style
            b_elm = parse_xml(f'<w:{border_name} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="{space}" w:color="{color}"/>')
            borders_elm.append(b_elm)
        else:
            b_elm = parse_xml(f'<w:{border_name} {nsdecls("w")} w:val="none"/>')
            borders_elm.append(b_elm)
    tc_pr.append(borders_elm)


def format_paragraph(p, space_before=0, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    p.alignment = align


def add_p(doc, text="", bold=False, italic=False, color=COLOR_TEXT_DARK, size=10.5,
          space_before=0, space_after=4, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    format_paragraph(p, space_before, space_after, line_spacing, align)
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.name = "Calibri"
        run.font.size = Pt(size)
        run.font.color.rgb = color
    return p


def add_heading_1(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=16, space_after=6, line_spacing=1.1, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(16.5)
    run.font.color.rgb = COLOR_PRIMARY
    return p


def add_heading_2(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=11, space_after=4, line_spacing=1.1, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.color.rgb = COLOR_SECONDARY
    return p


def add_heading_3(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, space_before=8, space_after=3, line_spacing=1.1, align=WD_ALIGN_PARAGRAPH.LEFT)
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Calibri"
    run.font.size = Pt(11)
    run.font.color.rgb = COLOR_TEXT_DARK
    return p


def add_bullet(doc, text, bold_prefix="", level=0):
    p = doc.add_paragraph(style='List Bullet')
    format_paragraph(p, space_before=1, space_after=3, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.bold = True
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(10.5)
        r_pre.font.color.rgb = COLOR_TEXT_DARK
    run = p.add_run(text)
    run.font.name = "Calibri"
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_TEXT_DARK
    return p


def add_callout(doc, title, text, icon="📌"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.8)
    set_cell_shading(cell, COLOR_LIGHT_BG_HEX)
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    set_cell_borders(cell, left=("single", "24", "0", COLOR_PRIMARY_HEX),
                           top=("single", "4", "0", COLOR_BORDER_HEX),
                           bottom=("single", "4", "0", COLOR_BORDER_HEX),
                           right=("single", "4", "0", COLOR_BORDER_HEX))
    
    p = cell.paragraphs[0]
    format_paragraph(p, space_before=0, space_after=3, line_spacing=1.15)
    run_title = p.add_run(f"{icon} {title}\n")
    run_title.bold = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(10.5)
    run_title.font.color.rgb = COLOR_PRIMARY
    
    run_text = p.add_run(text)
    run_text.font.name = "Calibri"
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = COLOR_TEXT_DARK
    
    p_spacer = doc.add_paragraph()
    format_paragraph(p_spacer, space_before=0, space_after=3)


def add_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        format_paragraph(p, space_before=3, space_after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
        run = p.add_run(title)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(hdr_cells[i], COLOR_PRIMARY_HEX)
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        hdr_cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    # Data Rows
    for r_idx, row_values in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_color = COLOR_ALT_ROW_HEX if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_values):
            row_cells[c_idx].text = ""
            p = row_cells[c_idx].paragraphs[0]
            format_paragraph(p, space_before=2, space_after=2, align=WD_ALIGN_PARAGRAPH.LEFT)
            run = p.add_run(str(val))
            run.font.name = "Calibri"
            run.font.size = Pt(9.5)
            run.font.color.rgb = COLOR_TEXT_DARK
            if c_idx == 0 and len(headers) > 2:
                run.bold = True
            set_cell_shading(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=100, right=100)
            set_cell_borders(row_cells[c_idx],
                             top=("single", "4", "0", COLOR_BORDER_HEX),
                             bottom=("single", "4", "0", COLOR_BORDER_HEX),
                             left=("single", "4", "0", COLOR_BORDER_HEX),
                             right=("single", "4", "0", COLOR_BORDER_HEX))
            row_cells[c_idx].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    if col_widths:
        for row in table.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)

    p_spacer = doc.add_paragraph()
    format_paragraph(p_spacer, space_before=0, space_after=3)
    return table


def add_image_with_caption(doc, image_filename, figure_number, title, desc_paragraphs):
    image_path = IMAGES_DIR / image_filename
    if not image_path.exists():
        print(f"WARNING: Image not found: {image_path}")
        return

    with Image.open(image_path) as im:
        orig_w, orig_h = im.size
        aspect_ratio = orig_h / orig_w

    if aspect_ratio > 1.15:       # Tall screenshot
        display_width = Inches(4.5)
    elif aspect_ratio >= 0.85:    # Square screenshot
        display_width = Inches(5.1)
    elif aspect_ratio < 0.45:     # Wide/narrow banner screenshot
        display_width = Inches(6.6)
    else:                         # Standard landscape screenshot
        display_width = Inches(6.4)

    # Image paragraph
    p_img = doc.add_paragraph()
    format_paragraph(p_img, space_before=8, space_after=3, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_img.paragraph_format.keep_with_next = True
    p_img.add_run().add_picture(str(image_path), width=display_width)

    # Caption paragraph
    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, space_before=2, space_after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
    p_cap.paragraph_format.keep_with_next = True
    run_num = p_cap.add_run(f"Figure {figure_number}: ")
    run_num.bold = True
    run_num.font.name = "Calibri"
    run_num.font.size = Pt(9.5)
    run_num.font.color.rgb = COLOR_PRIMARY

    run_title = p_cap.add_run(title)
    run_title.italic = True
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(9.5)
    run_title.font.color.rgb = COLOR_MUTED

    # Explanatory paragraphs
    for desc in desc_paragraphs:
        if isinstance(desc, tuple):
            prefix, body = desc
            p_desc = doc.add_paragraph()
            format_paragraph(p_desc, space_before=2, space_after=4, line_spacing=1.15)
            r_pre = p_desc.add_run(prefix)
            r_pre.bold = True
            r_pre.font.name = "Calibri"
            r_pre.font.size = Pt(10)
            r_pre.font.color.rgb = COLOR_SECONDARY
            r_body = p_desc.add_run(body)
            r_body.font.name = "Calibri"
            r_body.font.size = Pt(10)
            r_body.font.color.rgb = COLOR_TEXT_DARK
        else:
            add_p(doc, desc, size=10, space_before=2, space_after=4)

    p_spacer = doc.add_paragraph()
    format_paragraph(p_spacer, space_before=0, space_after=4)


def setup_document():
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

        footer = section.footer.paragraphs[0]
        format_paragraph(footer, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.RIGHT)
        r_ft = footer.add_run("AgriYield AI — Final Project Report | OVS VARUN")
        r_ft.font.name = "Calibri"
        r_ft.font.size = Pt(8.5)
        r_ft.font.color.rgb = COLOR_MUTED

        header = section.header.paragraphs[0]
        format_paragraph(header, space_before=0, space_after=0, align=WD_ALIGN_PARAGRAPH.LEFT)
        r_hd = header.add_run("AgriYield AI: AI-Powered Crop Yield Forecasting & Decision Support System")
        r_hd.font.name = "Calibri"
        r_hd.font.size = Pt(8.5)
        r_hd.font.color.rgb = COLOR_MUTED

    doc.core_properties.title = "AgriYield AI — Final Project Report"
    doc.core_properties.author = "OVS VARUN"
    doc.core_properties.comments = "Comprehensive Final Project Report for AgriYield AI Platform"
    return doc

