# -*- coding: utf-8 -*-
"""
AgriYield AI Report Generator - Chapter 7: Administrator Experience & Governance Controls
Contains Figures 18 to 21 with exhaustive visual, backend, and governance explanations.
"""

from docx.shared import Inches, Pt
from .common import (
    add_p, add_heading_1, add_heading_2, add_heading_3,
    add_bullet, add_callout, add_table, add_image_with_caption,
    COLOR_PRIMARY, COLOR_SECONDARY, COLOR_ACCENT
)


def write_chapter_7(doc):
    add_heading_1(doc, "7. Administrator Experience & Governance Controls")

    add_p(doc,
          "The Administrator workspace in AgriYield AI is designed as a centralized platform command center. It provides platform-level intelligence, user lifecycle governance, global prediction telemetry, security audit tracking, and dataset management without cluttering the farmer experience. Figures 18 through 21 document each administrative subsystem.",
          size=10.5)

    # ----------------------------------------------------
    # FIGURE 18: ADMIN DASHBOARD
    # ----------------------------------------------------
    add_heading_2(doc, "7.1 Administrator Platform Intelligence & Operations Dashboard")
    add_image_with_caption(
        doc,
        image_filename="admin_dashboard.png",
        figure_number=18,
        title="Administrator Intelligence Command Center with Platform KPIs and Crop Analytics",
        desc_paragraphs=[
            ("Interface Architecture & Visual Elements: ",
             "Figure 18 illustrates the Administrator Command Center. The top metric strip presents four platform KPIs: Total Registered Users (with farmer vs. admin breakdown), Total Predictions Executed, Average Forecasted Yield across all sessions (Kg/ha), and Active User Sessions. Below the KPI cards, interactive graphical widgets showcase: Most Frequently Forecasted Crops (bar chart showing Rice, Wheat, Cotton, Maize, and Groundnut distribution), Top Active States & Districts (geographic activity map), and Recent System Activity Feed. An administrative action bar enables one-click broadcast announcements, database backups, and global report exports."),
            ("Backend Processing & Microservices Flow: ",
             "When an administrator accesses `/admin`, the React application calls `GET /api/admin/stats`. The FastAPI backend verifies that the caller's JWT token contains `role: 'admin'`. If a non-admin attempts to access this endpoint, FastAPI returns an HTTP 403 Forbidden exception. The backend computes platform aggregates by reading `data/users.json` and `data/predictions.json`, calculating real-time frequency distributions, active account tallies, and mean yields in memory within 40 milliseconds."),
            ("Governance Value & Administrative Operations: ",
             "The command center gives platform managers real-time visibility into agricultural demand and user engagement. Administrators can identify which regions or crops are experiencing peak advisory demand, allowing extension agencies and agricultural cooperatives to align physical resources (such as fertilizer stock or hybrid seed supplies) with actual farmer interest.")
        ]
    )

    # ----------------------------------------------------
    # FIGURE 19: USER MANAGEMENT
    # ----------------------------------------------------
    add_heading_2(doc, "7.2 User Lifecycle Administration & Role-Based Access Control")
    add_image_with_caption(
        doc,
        image_filename="user_management.png",
        figure_number=19,
        title="User Lifecycle Management Table with Role Assignment and Status Controls",
        desc_paragraphs=[
            ("Interface Architecture & Visual Elements: ",
             "Figure 19 displays the User Management Administration panel. The interface consists of a responsive data table listing all registered accounts with columns: User UUID, Full Name, Email Address, Role (Farmer / Admin), Account Status (Active / Pending / Disabled), Registration Date, and an Action Menu. A global search bar enables instant filtering by name or email, while role and status filter chips narrow down account lists. Action buttons allow administrators to toggle account status (Activate, Suspend, Disable), reassign user roles, or reset account lockout states."),
            ("Backend Processing & Microservices Flow: ",
             "Account modifications trigger `PATCH /api/admin/users/{user_id}`. The FastAPI endpoint enforces strict security constraints: an administrator cannot delete or disable their own active account, preventing accidental platform lockout. Role promotions or status suspensions are immediately saved to `data/users.json`, and an immutable event record is created in the administrative audit log detailing who modified the account, the prior state, the new state, and the timestamp."),
            ("Governance Value & Administrative Operations: ",
             "User management provides granular access control and fraud mitigation. If an account exhibits abusive behavior, abnormal API traffic, or credential compromise, administrators can immediately suspend the account, preventing malicious access to platform prediction microservices while preserving historical audit records.")
        ]
    )

    # ----------------------------------------------------
    # FIGURE 20: PREDICTION MONITORING
    # ----------------------------------------------------
    add_heading_2(doc, "7.3 Platform-Wide Prediction Activity Telemetry & Monitoring")
    add_image_with_caption(
        doc,
        image_filename="prediction_monitoring.png",
        figure_number=20,
        title="Global Prediction Monitoring Table with Filtering and CSV Report Export",
        desc_paragraphs=[
            ("Interface Architecture & Visual Elements: ",
             "Figure 20 illustrates the Global Prediction Activity Monitoring interface. This screen provides an exhaustive, platform-wide record of every yield forecast generated across all farmers. The data table displays: Prediction ID, Farmer Name & Email, State & District, Crop Variety, Target Harvest Year, Cultivated Area (ha), Forecasted Yield (Kg/ha), Total Production (MT), AI Provider Mode utilized, and Timestamp. Comprehensive multi-criteria filters enable administrators to query predictions by date range, specific crop, state, or area threshold. A prominent 'Export CSV Report' button downloads the filtered dataset."),
            ("Backend Processing & Microservices Flow: ",
             "The table queries `GET /api/admin/predictions` with optional query parameters (`crop`, `state`, `year`, `limit`, `offset`). The FastAPI handler applies server-side filtering across `data/predictions.json`. For CSV exports, the backend streams the filtered dataset formatted in RFC 4180 compliant CSV format, complete with appropriate `Content-Disposition: attachment; filename=agri_predictions_report.csv` response headers."),
            ("Governance Value & Administrative Operations: ",
             "Global prediction monitoring transforms individual farmer queries into macro-level agricultural intelligence. District agricultural officers and policy makers can analyze aggregated crop demand, forecast regional grain surpluses or deficits, and plan regional grain procurement and storage infrastructure well in advance of the harvest season.")
        ]
    )

    # ----------------------------------------------------
    # FIGURE 21: AUDIT TRAIL LOGS
    # ----------------------------------------------------
    add_heading_2(doc, "7.4 Security Audit Trail & Authentication Event Logging")
    add_image_with_caption(
        doc,
        image_filename="authentication_login_audit_trail_logs.png",
        figure_number=21,
        title="Security Audit Trail Displaying Login Events, IP Addresses, and Lockouts",
        desc_paragraphs=[
            ("Interface Architecture & Visual Elements: ",
             "Figure 21 showcases the Security Audit Trail and Authentication Log interface. The table provides an immutable record of security and administrative operations across the platform. Columns detail: Event ID, Event Type (Login Success, Login Failure, Account Lockout, User Role Modified, Announcement Published, Dataset Version Registered), Initiator Email & Role, Client IP Address, User Agent / Browser Details, Timestamp (ISO-8601), and Event Outcome Status (Success / Warning / Failure). Color-coded badges highlight security-sensitive events, such as brute-force lockout triggers."),
            ("Backend Processing & Microservices Flow: ",
             "Every authentication attempt, role alteration, and administrative configuration change triggers an internal event dispatch to `backend/admin_operations.py`. The audit logger writes an append-only JSON record to `data/audit_log.json` and `data/login_history.json`. These files operate under thread-safe file locking to prevent race conditions during concurrent access. The admin UI fetches these records via `GET /api/admin/audit-log` and `GET /api/admin/login-history`."),
            ("Governance Value & Administrative Operations: ",
             "Comprehensive audit trails are essential for regulatory compliance, data integrity, and cyber-security forensics. If an unauthorized attempt is made to breach an administrative account or scrape prediction endpoints, the audit log captures the attacker's IP address, timestamp, and attempt frequency, allowing administrators to implement network-level blocks or review security policies.")
        ]
    )

    doc.add_page_break()
