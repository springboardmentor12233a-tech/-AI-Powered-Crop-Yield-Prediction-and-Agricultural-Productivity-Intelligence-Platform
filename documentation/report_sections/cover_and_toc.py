# -*- coding: utf-8 -*-
"""
AgriYield AI Report Generator - Cover Page, Table of Contents, and List of Figures
"""

from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from .common import (
    add_p, add_heading_1, add_callout, add_table, format_paragraph,
    COLOR_PRIMARY, COLOR_SECONDARY, COLOR_ACCENT, COLOR_MUTED
)


def add_cover_page(doc):
    for _ in range(3):
        p_space = doc.add_paragraph()
        format_paragraph(p_space, space_before=0, space_after=0)

    p_tag = doc.add_paragraph()
    format_paragraph(p_tag, space_before=0, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_tag = p_tag.add_run("MAJOR PROJECT FINAL REPORT | ARTIFICIAL INTELLIGENCE IN AGRICULTURE")
    r_tag.bold = True
    r_tag.font.name = "Calibri"
    r_tag.font.size = Pt(11)
    r_tag.font.color.rgb = COLOR_ACCENT

    p_title = doc.add_paragraph()
    format_paragraph(p_title, space_before=6, space_after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_title = p_title.add_run("AgriYield AI")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(36)
    r_title.font.color.rgb = COLOR_PRIMARY

    p_sub = doc.add_paragraph()
    format_paragraph(p_sub, space_before=4, space_after=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_sub = p_sub.add_run("AI-Powered Crop Yield Prediction & Agricultural Decision Support Platform")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(15.5)
    r_sub.font.color.rgb = COLOR_SECONDARY

    p_div = doc.add_paragraph()
    format_paragraph(p_div, space_before=0, space_after=16, align=WD_ALIGN_PARAGRAPH.CENTER)
    r_div = p_div.add_run("━" * 46)
    r_div.font.color.rgb = COLOR_PRIMARY
    r_div.bold = True

    meta_headers = ["Project Specification", "Submission Details"]
    meta_data = [
        ["Project Title", "AgriYield AI — Integrated Agricultural Intelligence System"],
        ["Candidate / Author", "OVS VARUN"],
        ["Candidate Email", "orugantivarun.2007@gmail.com"],
        ["Domain / Specialization", "Applied Machine Learning, Agro-Informatics & Cloud Engineering"],
        ["Technology Stack", "React 19, Vite, FastAPI, Python 3.12, XGBoost, Scikit-learn, Gemini 1.5, Groq"],
        ["Core ML Model", "Extreme Gradient Boosting (XGBRegressor) — R² = 0.8540, MAE = 237.72 Kg/ha"],
        ["Empirical Dataset", "ICRISAT Indian Meso-Level District Agricultural Dataset (12,418 records, 311 districts)"],
        ["Target Beneficiaries", "Smallholder Farmers, Agronomists, Cooperatives & Agricultural Administrators"],
        ["Project Milestone & Status", "Final Integrated Project Submission & Complete System Verification"],
        ["Submission Date", "October 2026"]
    ]
    add_table(doc, meta_headers, meta_data, [2.2, 4.6])

    add_callout(doc, "Project Purpose & Scope",
                "This comprehensive final project report documents the design, empirical machine-learning formulation, backend microservices architecture, frontend user experience, role-based governance controls, end-to-end testing verification, and deployment readiness of AgriYield AI. The system transforms decades of district-level agricultural data into actionable, high-precision yield forecasts paired with localized agronomic recommendations, ensuring data-driven farm empowerment across India.",
                icon="🌾")

    doc.add_page_break()


def add_table_of_contents(doc):
    add_heading_1(doc, "Table of Contents")
    toc_items = [
        ("1. Executive Summary & Project Abstract", "Comprehensive overview of AgriYield AI, problem context, and delivered solutions."),
        ("2. Problem Statement, Research Background & Objectives", "Challenges in Indian farming, decision risks, and concrete system objectives."),
        ("3. System Architecture & Technology Stack", "Multi-tier architecture: Presentation, REST API, ML Engine, Generative AI, and Security."),
        ("4. Dataset Engineering & Machine Learning Formulation", "ICRISAT district agricultural dataset, feature engineering, lag metrics, and XGBoost evaluation."),
        ("5. System Modules & Functional Breakdown", "Ten functional subsystems powering farmer enablement and administrative oversight."),
        ("6. Farmer Experience & Screen Walkthrough (Figures 1 – 17)", "Deep visual and technical walkthrough of public landing, authentication, dashboard, predictor, and tools."),
        ("7. Administrator Experience & Governance Controls (Figures 18 – 21)", "Command center, user administration, global prediction monitoring, and authentication audit logs."),
        ("8. Backend REST API Architecture & Endpoints", "Comprehensive specification of FastAPI endpoints, schemas, authentication, and HTTP responses."),
        ("9. System Testing, Verification & API Evidence (Figures 22 – 24)", "End-to-end automated testing, endpoint verification evidence, security lockout tests, and build validation."),
        ("10. Security Controls, Reliability & Performance Engineering", "Bcrypt hashing, JWT tokens, 5-strike brute-force lockout, role authorization, and fallback resilience."),
        ("11. Results, Agronomic Impact & Usability Analysis", "Empirical prediction outcomes, farmer workflow efficiency, and agricultural economic benefits."),
        ("12. Limitations & Edge Cases", "Granularity limitations, database concurrency scope, and external AI dependencies."),
        ("13. Future Roadmap & Commercialization", "PostgreSQL migration, satellite NDVI telemetry, IoT sensor integration, and mobile apps."),
        ("14. Conclusion", "Final synthesis of project achievements and closing remarks."),
        ("15. References & Data Sources", "Formal academic, technical, and data citations.")
    ]
    for title, desc in toc_items:
        p = doc.add_paragraph()
        format_paragraph(p, space_before=2, space_after=2)
        r_t = p.add_run(f"• {title}\n  ")
        r_t.bold = True
        r_t.font.name = "Calibri"
        r_t.font.size = Pt(10.5)
        r_t.font.color.rgb = COLOR_PRIMARY
        r_d = p.add_run(desc)
        r_d.italic = True
        r_d.font.name = "Calibri"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_MUTED

    doc.add_page_break()

    add_heading_1(doc, "List of Figures & Screen Evidence")
    add_p(doc, "This report contains 24 high-resolution figures representing actual screens, workflows, administrative controls, and API testing evidence captured from the operational AgriYield AI system:", italic=True, size=10)
    
    figures_list = [
        ("Figure 1", "Public Landing Page & Unified Agricultural Portal", "Images/main_page.png"),
        ("Figure 2", "Role-Aware Authentication & Sign-In Interface", "Images/signin_farmer_admin.png"),
        ("Figure 3", "Farmer & Administrator Registration System", "Images/Register_farmer_admin.png"),
        ("Figure 4", "Farmer Operational Command Center & Agricultural Overview", "Images/farmer_dashboard.png"),
        ("Figure 5", "Multilingual Localization & Regional Language Accessibility", "Images/lanuages.png"),
        ("Figure 6", "Crop Yield Prediction Input Form & AI Mode Selector", "Images/yield_predictor_1.png"),
        ("Figure 7", "Machine Learning Yield Forecast Output & Agronomic Advisory", "Images/yield_predictor_2.png"),
        ("Figure 8", "Soil NPK Nutrient Balance & Crop-Specific Fertilizer Advisory", "Images/soil_npk.png"),
        ("Figure 9", "Comprehensive Soil Health, Moisture & Deficiencies Evaluation", "Images/soil_analysis.png"),
        ("Figure 10", "Multi-Crop Comparative Analysis & Economic Planning", "Images/multi_crop.png"),
        ("Figure 11", "Historical Prediction Telemetry & Farmer Activity Records", "Images/farmer_logs.png"),
        ("Figure 12", "AI Agronomist Interactive Chat & Contextual Advisory Assistant", "Images/ai_assistant.png"),
        ("Figure 13", "Farmer Utility Suite & Specialized Agronomic Tools Hub", "Images/farmer_tools.png"),
        ("Figure 14", "Agro-Meteorological Weather Monitoring & Microclimate Analysis", "Images/weather_analysis.png"),
        ("Figure 15", "Multi-Season Crop Rotation & Soil Fertility Replenishment Planner", "Images/crop_rotation.png"),
        ("Figure 16", "Precision Seed Quantity & Sowing Density Calculator", "Images/seed_quantity_calcultor.png"),
        ("Figure 17", "Integrated Pest & Disease Diagnostic Decision System", "Images/pest_disease_diagonsis.png"),
        ("Figure 18", "Administrator Platform Intelligence & Operations Dashboard", "Images/admin_dashboard.png"),
        ("Figure 19", "User Lifecycle Administration & Role-Based Access Control", "Images/user_management.png"),
        ("Figure 20", "Platform-Wide Prediction Activity Telemetry & Monitoring", "Images/prediction_monitoring.png"),
        ("Figure 21", "Security Audit Trail & Authentication Event Logging", "Images/authentication_login_audit_trail_logs.png"),
        ("Figure 22", "Automated API Verification: User Registration Endpoint", "Images/api_testing_register_endpoint.png"),
        ("Figure 23", "Automated API Verification: OAuth2/JWT Authentication Login Endpoint", "Images/api_testing_login_endpoint.png"),
        ("Figure 24", "Automated API Verification: Machine Learning Inference Endpoint", "Images/api_testing_predict_endpoint.png")
    ]
    fig_table_data = [[fig_num, title, file_path] for fig_num, title, file_path in figures_list]
    add_table(doc, ["Figure #", "Visual Caption & Screen Title", "Source Image Asset"], fig_table_data, [1.1, 3.8, 1.9])
    
    doc.add_page_break()
